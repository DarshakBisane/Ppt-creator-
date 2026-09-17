"""Unit and integration tests for GenerationOrchestrator."""

import io
from pathlib import Path
import pytest
import pptx
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from backend.app.ai.context import PresentationGenerationRequest
from backend.app.ai.fake_provider import FakeAIProvider
from backend.app.artifacts.storage import ArtifactStorage
from backend.app.domain.enums import JobStatus
from backend.app.jobs.orchestrator import GenerationOrchestrator
from backend.app.jobs.store import MemoryJobStore


def create_sample_reference_pptx() -> bytes:
    """Helper generating synthetic dark-theme reference PPTX bytes."""
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(15, 23, 42)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf.read()


@pytest.mark.asyncio
async def test_mode_a_generation_orchestration(tmp_path: Path) -> None:
    """Test full Mode A background generation pipeline to completion."""
    store = MemoryJobStore()
    artifacts = ArtifactStorage(base_dir=tmp_path)
    ai_provider = FakeAIProvider()
    orchestrator = GenerationOrchestrator(
        job_store=store,
        artifact_storage=artifacts,
        ai_provider=ai_provider,
    )

    req = PresentationGenerationRequest(
        mode="topic",
        topic="Cloud Architecture & Kubernetes",
        slide_count=5,
        style="modern_dark",
    )
    store.create_job("job_mode_a", req)

    await orchestrator.execute_generation_job("job_mode_a")

    job = store.get_job("job_mode_a")
    assert job is not None
    assert job.status == JobStatus.COMPLETED
    assert job.progress.percent == 100

    resp = store.get_status_response("job_mode_a")
    assert resp is not None
    assert resp.state == "completed"
    assert resp.artifact is not None
    assert resp.artifact.download_url == "/api/download/job_mode_a"

    # Verify artifact on disk
    artifact_data = artifacts.get_artifact("job_mode_a")
    assert artifact_data is not None
    data, filename, path = artifact_data
    assert len(data) > 3000
    assert filename.endswith(".pptx")


@pytest.mark.asyncio
async def test_mode_b_reference_generation_orchestration(tmp_path: Path) -> None:
    """Test Mode B reference PPTX style transfer generation pipeline."""
    store = MemoryJobStore()
    artifacts = ArtifactStorage(base_dir=tmp_path)
    ai_provider = FakeAIProvider()
    orchestrator = GenerationOrchestrator(
        job_store=store,
        artifact_storage=artifacts,
        ai_provider=ai_provider,
    )

    ref_bytes = create_sample_reference_pptx()
    req = PresentationGenerationRequest(
        mode="reference",
        topic="Adaptive Style Transfer Architecture",
        slide_count=4,
    )
    store.create_job("job_mode_b", req, reference_bytes=ref_bytes)

    await orchestrator.execute_generation_job("job_mode_b")

    job = store.get_job("job_mode_b")
    assert job is not None
    assert job.status == JobStatus.COMPLETED
    assert job.progress.percent == 100

    artifact_data = artifacts.get_artifact("job_mode_b")
    assert artifact_data is not None


@pytest.mark.asyncio
async def test_orchestrator_failure_handling(tmp_path: Path) -> None:
    """Test that errors in the pipeline mark the job failed and clean artifacts."""
    store = MemoryJobStore()
    artifacts = ArtifactStorage(base_dir=tmp_path)

    class FailingAIProvider:
        async def generate_presentation(self, request: PresentationGenerationRequest):
            raise RuntimeError("Synthetic AI engine connection failure")

    orchestrator = GenerationOrchestrator(
        job_store=store,
        artifact_storage=artifacts,
        ai_provider=FailingAIProvider(),  # type: ignore
    )

    req = PresentationGenerationRequest(topic="Will Fail", slide_count=3)
    store.create_job("job_failing", req)

    await orchestrator.execute_generation_job("job_failing")

    job = store.get_job("job_failing")
    assert job is not None
    assert job.status == JobStatus.FAILED
    assert job.error is not None
    assert job.error.code == "AI_PROVIDER_ERROR"
    assert "AI presentation planning failed" in job.error.message
    # No partial artifact remains
    assert artifacts.get_artifact("job_failing") is None
