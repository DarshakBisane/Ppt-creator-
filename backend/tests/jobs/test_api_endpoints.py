"""Integration tests for FastAPI generation, status, and download endpoints."""

import io
from fastapi.testclient import TestClient
import pytest

from backend.app.artifacts.storage import default_artifact_storage
from backend.app.config import Settings
from backend.app.domain.enums import JobStatus
from backend.app.jobs.store import default_job_store
from backend.app.main import create_app


@pytest.fixture
def api_client(tmp_path) -> TestClient:
    """Fixture providing TestClient with test settings."""
    settings = Settings(
        app_name="Test Presentation API",
        environment="test",
        artifact_dir=str(tmp_path),
        log_level="DEBUG",
    )
    # Set artifact storage base dir
    default_artifact_storage.base_dir = tmp_path
    app = create_app(settings=settings)
    return TestClient(app, base_url="http://testserver")


def test_post_generate_json_mode_a(api_client: TestClient) -> None:
    """Test POST /api/generate with JSON body creates queued job."""
    payload = {
        "mode": "topic",
        "topic": "Microservices with Rust and Tokio",
        "audience": "engineers",
        "purpose": "technical",
        "slide_count": 5,
        "style": "modern_dark",
    }
    response = api_client.post("/api/generate", json=payload)
    assert response.status_code == 202
    data = response.json()
    assert "job_id" in data
    assert data["state"] == "queued"
    assert data["status_url"].startswith("/api/jobs/")

    job_id = data["job_id"]
    # Poll job status
    status_resp = api_client.get(f"/api/jobs/{job_id}")
    assert status_resp.status_code == 200
    status_data = status_resp.json()
    assert status_data["job_id"] == job_id


def test_post_generate_multipart_mode_b(api_client: TestClient) -> None:
    """Test POST /api/generate with multipart form-data and reference file upload."""
    sample_pptx_data = b"PK\x03\x04" + b"\x00" * 100  # Dummy header
    files = {
        "reference_file": ("template.pptx", sample_pptx_data, "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    }
    data = {
        "topic": "Adaptive Style Transfer",
        "mode": "reference",
        "slide_count": "6",
        "style": "professional",
    }

    response = api_client.post("/api/generate", data=data, files=files)
    assert response.status_code == 202
    res_json = response.json()
    assert "job_id" in res_json


def test_get_job_status_404_for_missing_job(api_client: TestClient) -> None:
    """Test GET /api/jobs/{job_id} returns 404 for unknown job ID."""
    response = api_client.get("/api/jobs/non_existent_job_123")
    assert response.status_code == 404
    err_data = response.json()
    assert err_data["error"]["code"] == "JOB_NOT_FOUND"


def test_download_completed_presentation(api_client: TestClient) -> None:
    """Test GET /api/download/{job_id} serves valid PPTX when completed."""
    # Create minimal PPTX in artifact storage
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("[Content_Types].xml", b"<Types></Types>")
        zf.writestr("ppt/presentation.xml", b"<p:presentation></p:presentation>")
    pptx_bytes = buf.getvalue()

    job_id = "test_download_job"
    default_artifact_storage.save_artifact(job_id, "Cloud_Strategy.pptx", pptx_bytes)
    default_job_store.create_job(job_id, None)  # type: ignore
    default_job_store.update_progress(job_id, "planning", 15, "Planning", JobStatus.PLANNING)
    default_job_store.update_progress(job_id, "designing", 35, "Designing", JobStatus.DESIGNING)
    default_job_store.update_progress(job_id, "qa", 75, "QA", JobStatus.QA)
    default_job_store.update_progress(job_id, "rendering", 90, "Rendering", JobStatus.RENDERING)
    default_job_store.complete_job(job_id, "path", "Cloud_Strategy.pptx", 5, len(pptx_bytes))

    response = api_client.get(f"/api/download/{job_id}")
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert "Cloud_Strategy.pptx" in response.headers["content-disposition"]
    assert response.content == pptx_bytes


def test_download_404_when_artifact_missing(api_client: TestClient) -> None:
    """Test GET /api/download/{job_id} returns 404 when artifact does not exist."""
    response = api_client.get("/api/download/missing_artifact_job")
    assert response.status_code == 404


def test_path_traversal_rejection_on_download(api_client: TestClient) -> None:
    """Test that path traversal attempts in download endpoint are rejected."""
    response = api_client.get("/api/download/..%2F..%2Fetc%2Fpasswd")
    assert response.status_code in (404, 400)


def test_path_traversal_rejection_on_job_status(api_client: TestClient) -> None:
    """Test that path traversal attempts in job status endpoint are rejected."""
    response = api_client.get("/api/jobs/..%2F..%2Fetc%2Fpasswd")
    assert response.status_code in (404, 400)


def test_post_generate_validation_error_missing_topic(api_client: TestClient) -> None:
    """Test POST /api/generate without topic fails with 422."""
    response = api_client.post("/api/generate", json={"mode": "topic", "topic": ""})
    assert response.status_code == 422
    err_data = response.json()
    assert err_data["error"]["code"] == "VALIDATION_ERROR"


def test_e2e_api_mode_a_execution_and_download(api_client: TestClient) -> None:
    """Test complete end-to-end Mode A execution via API client."""
    import zipfile
    from pptx import Presentation as PPTXDoc

    # 1. Submit Generation
    post_resp = api_client.post(
        "/api/generate",
        json={
            "mode": "topic",
            "topic": "Autonomous AI Agent Swarms",
            "audience": "technical",
            "purpose": "architecture",
            "slide_count": 4,
            "style": "modern_dark",
        },
    )
    assert post_resp.status_code == 202
    job_id = post_resp.json()["job_id"]

    # 2. Poll Status (FastAPI BackgroundTasks run synchronously during TestClient requests)
    status_resp = api_client.get(f"/api/jobs/{job_id}")
    assert status_resp.status_code == 200
    status_data = status_resp.json()
    assert status_data["state"] == "completed"
    assert status_data["progress"] == 100
    assert status_data["stage"] == "completed"
    assert status_data["artifact"] is not None
    assert status_data["artifact"]["download_url"] == f"/api/download/{job_id}"

    # 3. Download Artifact
    dl_resp = api_client.get(f"/api/download/{job_id}")
    assert dl_resp.status_code == 200
    assert dl_resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    assert len(dl_resp.content) > 0

    # 4. Verify OpenXML Package
    pptx_io = io.BytesIO(dl_resp.content)
    assert zipfile.is_zipfile(pptx_io)
    with zipfile.ZipFile(pptx_io, "r") as zf:
        names = zf.namelist()
        assert "[Content_Types].xml" in names
        assert "ppt/presentation.xml" in names
        assert any(n.startswith("ppt/slides/slide") for n in names)

    # 5. Verify Editable PPTX Content
    pptx_io.seek(0)
    prs = PPTXDoc(pptx_io)
    assert len(prs.slides) == 4
    # Confirm slide shapes exist and are editable native objects
    first_slide = prs.slides[0]
    assert len(first_slide.shapes) > 0


def test_e2e_api_mode_b_execution_and_download(api_client: TestClient) -> None:
    """Test complete end-to-end Mode B execution with reference file."""
    import zipfile
    from pptx import Presentation as PPTXDoc

    # Create dummy reference PPTX
    ref_prs = PPTXDoc()
    ref_io = io.BytesIO()
    ref_prs.save(ref_io)
    ref_bytes = ref_io.getvalue()

    files = {
        "reference_file": ("brand_template.pptx", ref_bytes, "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    }
    data = {
        "topic": "Distributed Consensus in Cloud Native Systems",
        "mode": "reference",
        "audience": "executive",
        "purpose": "strategy",
        "slide_count": 4,
        "style": "executive",
    }

    # 1. Submit Generation
    post_resp = api_client.post("/api/generate", data=data, files=files)
    assert post_resp.status_code == 202
    job_id = post_resp.json()["job_id"]

    # 2. Poll Status
    status_resp = api_client.get(f"/api/jobs/{job_id}")
    assert status_resp.status_code == 200
    status_data = status_resp.json()
    assert status_data["state"] == "completed"
    assert status_data["progress"] == 100
    assert status_data["artifact"] is not None

    # 3. Download and verify
    dl_resp = api_client.get(f"/api/download/{job_id}")
    assert dl_resp.status_code == 200
    assert len(dl_resp.content) > 0
    prs = PPTXDoc(io.BytesIO(dl_resp.content))
    assert len(prs.slides) == 4

