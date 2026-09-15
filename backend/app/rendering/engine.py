"""Native PowerPoint Presentation Renderer executing deterministic layouts."""

import io
from typing import Any
import pptx
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from backend.app.domain.design_system import DesignSystem
from backend.app.domain.enums import ElementType
from backend.app.domain.presentation import Presentation
from backend.app.layout.constants import CANVAS_HEIGHT, CANVAS_WIDTH, virtual_to_emu
from backend.app.layout.engine import resolve_presentation_layout
from backend.app.layout.models import PresentationLayoutResult
from backend.app.rendering.charts import render_chart_element
from backend.app.rendering.connectors import render_connector
from backend.app.rendering.shapes import (
    hex_to_rgb,
    render_card_element,
    render_kpi_element,
    render_shape_element,
)
from backend.app.rendering.tables import render_table_element
from backend.app.rendering.text import render_text_element


class PPTXRenderer:
    """Production Native PPTX Renderer consuming resolved layout geometry."""

    def render(
        self,
        presentation: Presentation,
        layout_result: Any | None = None,
    ) -> pptx.Presentation:
        """Render a Presentation domain model and its resolved layout into a python-pptx Presentation."""
        if hasattr(layout_result, "layout") and layout_result.layout is not None:
            layout = layout_result.layout
        else:
            layout = layout_result or resolve_presentation_layout(presentation)
        ds = presentation.design_system or DesignSystem()

        prs = pptx.Presentation()

        # Set canonical 16:9 Widescreen slide dimensions
        prs.slide_width = virtual_to_emu(CANVAS_WIDTH)
        prs.slide_height = virtual_to_emu(CANVAS_HEIGHT)

        # Blank layout template
        blank_layout = prs.slide_layouts[6]

        for slide_model, slide_layout in zip(presentation.slides, layout.slides):
            slide_pptx = prs.slides.add_slide(blank_layout)
            shapes = slide_pptx.shapes

            # 1. Slide Canvas Background
            bg_shape = shapes.add_shape(
                MSO_SHAPE.RECTANGLE,
                0,
                0,
                prs.slide_width,
                prs.slide_height,
            )
            bg_shape.fill.solid()
            bg_shape.fill.fore_color.rgb = hex_to_rgb(ds.palette.background.value)
            bg_shape.line.fill.background()

            # 2. Render Elements
            for elem in slide_layout.elements:
                if elem.semantic_type in (ElementType.HEADING, ElementType.TEXT):
                    render_text_element(shapes, elem, ds)
                elif elem.semantic_type == ElementType.CARD:
                    render_card_element(shapes, elem, ds)
                elif elem.semantic_type == ElementType.KPI:
                    render_kpi_element(shapes, elem, ds)
                elif elem.semantic_type == ElementType.TABLE:
                    render_table_element(shapes, elem, ds)
                elif elem.semantic_type == ElementType.CHART:
                    render_chart_element(shapes, elem, ds)
                else:
                    render_shape_element(shapes, elem, ds)

            # 3. Render Connectors
            for conn in slide_layout.connectors:
                try:
                    render_connector(shapes, conn, ds)
                except Exception:
                    # Skip problematic individual connectors gracefully
                    continue

            # 4. Attach Speaker Notes
            if slide_model.speaker_notes:
                notes_slide = slide_pptx.notes_slide
                text_frame = notes_slide.notes_text_frame
                text_frame.text = slide_model.speaker_notes

        return prs

    def render_to_bytes(
        self,
        presentation: Presentation,
        layout_result: PresentationLayoutResult | None = None,
    ) -> bytes:
        """Render presentation and return raw binary PPTX byte payload."""
        prs = self.render(presentation, layout_result)
        buffer = io.BytesIO()
        prs.save(buffer)
        buffer.seek(0)
        return buffer.read()

    def render_to_file(
        self,
        presentation: Presentation,
        filepath: str,
        layout_result: PresentationLayoutResult | None = None,
    ) -> None:
        """Render presentation and save directly to filesystem path."""
        prs = self.render(presentation, layout_result)
        prs.save(filepath)


# Global default renderer instance
default_pptx_renderer = PPTXRenderer()


def render_presentation_to_pptx(
    presentation: Presentation,
    layout_result: PresentationLayoutResult | None = None,
) -> pptx.Presentation:
    """Convenience helper to render a Presentation into a pptx Presentation object."""
    return default_pptx_renderer.render(presentation, layout_result)


def render_presentation_to_bytes(
    presentation: Presentation,
    layout_result: PresentationLayoutResult | None = None,
) -> bytes:
    """Convenience helper to render a Presentation into PPTX bytes."""
    return default_pptx_renderer.render_to_bytes(presentation, layout_result)
