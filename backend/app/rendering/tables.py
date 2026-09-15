"""Native PowerPoint table rendering."""

from typing import Any
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from backend.app.domain.design_system import DesignSystem
from backend.app.domain.tables import TableData
from backend.app.layout.models import ElementGeometry
from backend.app.rendering.shapes import hex_to_rgb


def render_table_element(
    slide_shape_tree: Any,
    elem: ElementGeometry,
    ds: DesignSystem,
) -> Any:
    """Render a native PowerPoint table populated from TableData."""
    x_emu, y_emu, w_emu, h_emu = elem.rect.to_emu()
    table_data: TableData | None = elem.content_data if isinstance(elem.content_data, TableData) else None

    if not table_data or not table_data.columns or not table_data.rows:
        # Default placeholder table
        cols_count = 3
        rows_count = 4
        headers = ["Capability", "Traditional Method", "PresenAI Platform"]
        sample_rows = [
            ["Generation Time", "4-8 Hours", "< 3 Seconds"],
            ["Layout Engine", "Manual Drag & Drop", "Deterministic 1920x1080"],
            ["Output Quality", "Rasterized / Fragile", "100% Native OpenXML"],
        ]
    else:
        cols_count = len(table_data.columns)
        rows_count = len(table_data.rows) + (1 if table_data.has_header else 0)
        headers = [c.label for c in table_data.columns]
        sample_rows = [r.cells for r in table_data.rows]

    table_shape = slide_shape_tree.add_table(
        rows_count,
        cols_count,
        x_emu,
        y_emu,
        w_emu,
        h_emu,
    )
    table = table_shape.table

    # Format Header Row
    for col_idx, header_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = hex_to_rgb(ds.palette.primary.value)
        cell.margin_left = Inches(0.1)
        cell.margin_right = Inches(0.1)

        p = cell.text_frame.paragraphs[0]
        p.text = header_text
        p.font.name = ds.typography.heading.font_family
        p.font.size = Pt(ds.typography.label.font_size)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb("#FFFFFF")
        p.alignment = PP_ALIGN.LEFT

    # Format Data Rows
    for row_idx, row_cells in enumerate(sample_rows, start=1):
        bg_hex = ds.palette.surface.value if row_idx % 2 == 1 else ds.palette.background.value
        for col_idx, cell_value in enumerate(row_cells):
            if col_idx >= cols_count or row_idx >= rows_count:
                continue
            cell = table.cell(row_idx, col_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = hex_to_rgb(bg_hex)
            cell.margin_left = Inches(0.1)
            cell.margin_right = Inches(0.1)

            p = cell.text_frame.paragraphs[0]
            p.text = str(cell_value)
            p.font.name = ds.typography.body.font_family
            p.font.size = Pt(ds.typography.body.font_size - 1)
            p.font.color.rgb = hex_to_rgb(ds.palette.text_primary.value)
            p.alignment = PP_ALIGN.LEFT

    return table_shape
