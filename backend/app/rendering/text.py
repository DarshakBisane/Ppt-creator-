"""Native PowerPoint text frame rendering."""

from typing import Any
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from backend.app.domain.design_system import DesignSystem
from backend.app.domain.enums import Alignment, ElementType
from backend.app.layout.models import ElementGeometry
from backend.app.rendering.shapes import hex_to_rgb


def extract_text(data: Any) -> str:
    """Extract string text from arbitrary content data containers."""
    if data is None:
        return ""
    if isinstance(data, str):
        return data
    if isinstance(data, dict):
        return str(data.get("text") or data.get("title") or "")
    if hasattr(data, "text"):
        return str(data.text)
    if hasattr(data, "model_dump"):
        d = data.model_dump()
        return str(d.get("text") or d.get("title") or "")
    return str(data)


def render_text_element(
    slide_shape_tree: Any,
    elem: ElementGeometry,
    ds: DesignSystem,
) -> Any:
    """Render a native PowerPoint textbox with styling derived from DesignSystem."""
    x_emu, y_emu, w_emu, h_emu = elem.rect.to_emu()
    txBox = slide_shape_tree.add_textbox(x_emu, y_emu, w_emu, h_emu)
    tf = txBox.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    raw_text = extract_text(elem.content_data)

    p = tf.paragraphs[0]
    p.text = str(raw_text)

    # Typography & alignment determination
    if elem.role == "slide_title" or elem.semantic_type == ElementType.HEADING:
        p.font.name = ds.typography.title.font_family
        p.font.size = Pt(ds.typography.title.font_size)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb(ds.palette.text_primary.value)
    elif elem.role == "slide_subtitle":
        p.font.name = ds.typography.body.font_family
        p.font.size = Pt(ds.typography.body.font_size + 2)
        p.font.color.rgb = hex_to_rgb(ds.palette.text_secondary.value)
    else:
        p.font.name = ds.typography.body.font_family
        p.font.size = Pt(ds.typography.body.font_size)
        p.font.color.rgb = hex_to_rgb(ds.palette.text_primary.value)

    if elem.alignment == Alignment.CENTER:
        p.alignment = PP_ALIGN.CENTER
    elif elem.alignment == Alignment.RIGHT:
        p.alignment = PP_ALIGN.RIGHT
    else:
        p.alignment = PP_ALIGN.LEFT

    return txBox

