"""Helper utilities and shape rendering for native PowerPoint cards and containers."""

from typing import Any
import pptx
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

from backend.app.domain.design_system import DesignSystem
from backend.app.layout.models import ElementGeometry, Rect


def hex_to_rgb(hex_str: str) -> RGBColor:
    """Convert hex string (e.g. #6366F1 or 6366F1) to pptx RGBColor."""
    clean = hex_str.strip().lstrip("#")
    if len(clean) == 3:
        clean = "".join([c * 2 for c in clean])
    if len(clean) != 6:
        return RGBColor(100, 100, 100)
    r = int(clean[0:2], 16)
    g = int(clean[2:4], 16)
    b = int(clean[4:6], 16)
    return RGBColor(r, g, b)


def to_dict(data: Any) -> dict[str, Any]:
    """Convert content data container or Pydantic model to dictionary."""
    if data is None:
        return {}
    if isinstance(data, dict):
        return data
    if hasattr(data, "model_dump"):
        return data.model_dump()
    if hasattr(data, "__dict__"):
        return data.__dict__
    return {"text": str(data)}


def render_card_element(
    slide_shape_tree: Any,
    elem: ElementGeometry,
    ds: DesignSystem,
) -> Any:
    """Render a native PowerPoint rounded card with formatted title and body text."""
    x_emu, y_emu, w_emu, h_emu = elem.rect.to_emu()

    # Determine background and border colors based on theme and role
    is_hero = elem.style_hints.get("variant") == "hero_banner" or elem.importance == "primary"
    is_takeaway = elem.style_hints.get("variant") == "takeaway_banner" or elem.role == "takeaway"
    bg_hex = ds.palette.surface.value
    if is_takeaway:
        border_hex = ds.palette.accent.value if ds.palette.accent else ds.palette.primary.value
    elif is_hero:
        border_hex = ds.palette.primary.value
    else:
        border_hex = ds.palette.surface.value

    text_primary_rgb = hex_to_rgb(ds.palette.text_primary.value)
    text_secondary_rgb = hex_to_rgb(ds.palette.text_secondary.value)
    primary_rgb = hex_to_rgb(ds.palette.primary.value)
    accent_rgb = hex_to_rgb(ds.palette.accent.value if ds.palette.accent else ds.palette.primary.value)

    shape = slide_shape_tree.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x_emu,
        y_emu,
        w_emu,
        h_emu,
    )

    # Shape Styling
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_to_rgb(bg_hex)
    shape.line.color.rgb = hex_to_rgb(border_hex)
    shape.line.width = Pt(2.0 if is_takeaway else (1.5 if is_hero else 1.0))

    # Text Framing inside card
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.12 if is_takeaway else 0.2)
    tf.margin_bottom = Inches(0.12 if is_takeaway else 0.2)

    content = to_dict(elem.content_data)
    if content:
        # Card title / headline
        title_text = content.get("title") or content.get("headline") or content.get("phase") or ""
        body_text = content.get("body") or content.get("text") or content.get("description") or content.get("subtitle") or ""
        bullets = content.get("bullets") or content.get("items") or content.get("points") or content.get("deliverables") or []
        step_num = content.get("step_number")

        p = tf.paragraphs[0]
        if is_takeaway:
            p.text = f"KEY TAKEAWAY  •  {body_text}"
            p.font.name = ds.typography.body.font_family
            p.font.size = Pt(13)
            p.font.bold = False
            p.font.color.rgb = text_primary_rgb
        else:
            if step_num:
                p.text = f"{step_num}  {title_text}"
            else:
                p.text = title_text

            p.font.name = ds.typography.heading.font_family
            p.font.size = Pt(ds.typography.heading.font_size)
            p.font.bold = True
            p.font.color.rgb = primary_rgb if is_hero else text_primary_rgb

            if body_text:
                p_body = tf.add_paragraph()
                p_body.text = body_text
                p_body.font.name = ds.typography.body.font_family
                p_body.font.size = Pt(ds.typography.body.font_size)
                p_body.font.color.rgb = text_secondary_rgb
                p_body.space_before = Pt(6)

            for b_item in bullets:
                p_bullet = tf.add_paragraph()
                p_bullet.text = f"• {b_item}"
                p_bullet.font.name = ds.typography.body.font_family
                p_bullet.font.size = Pt(ds.typography.body.font_size - 1)
                p_bullet.font.color.rgb = text_secondary_rgb
                p_bullet.space_before = Pt(4)

    return shape


def render_kpi_element(
    slide_shape_tree: Any,
    elem: ElementGeometry,
    ds: DesignSystem,
) -> Any:
    """Render a dedicated hero KPI stat card with massive metric value and label."""
    x_emu, y_emu, w_emu, h_emu = elem.rect.to_emu()
    shape = slide_shape_tree.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        x_emu,
        y_emu,
        w_emu,
        h_emu,
    )

    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_to_rgb(ds.palette.surface.value)
    shape.line.color.rgb = hex_to_rgb(ds.palette.primary.value)
    shape.line.width = Pt(1.5)

    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.2)
    tf.margin_right = Inches(0.2)
    tf.margin_top = Inches(0.25)
    tf.margin_bottom = Inches(0.2)

    content = to_dict(elem.content_data)
    value = content.get("value") or content.get("metric") or "100%"
    unit = content.get("unit") or ""
    label = content.get("label") or "Performance Metric"
    context = content.get("context") or content.get("subtext") or ""
    trend = content.get("trend")

    # Paragraph 1: Metric Value + Unit
    p_val = tf.paragraphs[0]
    p_val.text = f"{value} {unit}".strip()
    p_val.font.name = ds.typography.display.font_family
    p_val.font.size = Pt(36)
    p_val.font.bold = True
    p_val.font.color.rgb = hex_to_rgb(ds.palette.accent.value if ds.palette.accent else ds.palette.primary.value)
    p_val.alignment = PP_ALIGN.CENTER

    # Paragraph 2: Metric Label
    p_lbl = tf.add_paragraph()
    p_lbl.text = label
    p_lbl.font.name = ds.typography.heading.font_family
    p_lbl.font.size = Pt(ds.typography.label.font_size + 2)
    p_lbl.font.bold = True
    p_lbl.font.color.rgb = hex_to_rgb(ds.palette.text_primary.value)
    p_lbl.alignment = PP_ALIGN.CENTER
    p_lbl.space_before = Pt(6)

    # Paragraph 3: Context / Trend
    if context or trend:
        p_ctx = tf.add_paragraph()
        trend_str = f"▲ " if trend == "up" else (f"▼ " if trend == "down" else "")
        p_ctx.text = f"{trend_str}{context}".strip()
        p_ctx.font.name = ds.typography.caption.font_family
        p_ctx.font.size = Pt(ds.typography.caption.font_size)
        p_ctx.font.color.rgb = hex_to_rgb(ds.palette.text_secondary.value)
        p_ctx.alignment = PP_ALIGN.CENTER
        p_ctx.space_before = Pt(4)

    return shape


def render_shape_element(
    slide_shape_tree: Any,
    elem: ElementGeometry,
    ds: DesignSystem,
) -> Any:
    """Render generic geometric shapes such as pills, nodes, spines, and badges."""
    x_emu, y_emu, w_emu, h_emu = elem.rect.to_emu()

    variant = elem.style_hints.get("variant", "rect")
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if variant == "pill" else (MSO_SHAPE.OVAL if variant == "circle" else MSO_SHAPE.RECTANGLE)

    shape = slide_shape_tree.add_shape(
        shape_type,
        x_emu,
        y_emu,
        w_emu,
        h_emu,
    )

    shape.fill.solid()
    if elem.importance == "accent":
        shape.fill.fore_color.rgb = hex_to_rgb(ds.palette.accent.value)
        shape.line.fill.background()
    elif elem.role == "spine":
        shape.fill.fore_color.rgb = hex_to_rgb(ds.palette.primary.value)
        shape.line.fill.background()
    else:
        shape.fill.fore_color.rgb = hex_to_rgb(ds.palette.surface.value)
        shape.line.color.rgb = hex_to_rgb(ds.palette.primary.value)

    # Optional text in badge
    content = to_dict(elem.content_data)
    badge_text = content.get("text") or elem.style_hints.get("text")
    if badge_text:
        tf = shape.text_frame
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.text = str(badge_text)
        p.font.name = ds.typography.caption.font_family
        p.font.size = Pt(ds.typography.caption.font_size)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb("#FFFFFF") if elem.importance == "accent" else hex_to_rgb(ds.palette.text_primary.value)
        p.alignment = PP_ALIGN.CENTER

    return shape
